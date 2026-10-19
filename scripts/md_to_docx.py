#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> Word with real footnotes (Pandoc), optional claim-graph image.

Usage:
  python md_to_docx.py --md analysis.md --out analysis.docx [--image graph.png] [--citations citations.json]

If --citations is given and the Markdown has no footnote marker, the
citation_law_footnote is injected as a real Word footnote near the beginning.
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys


def _inject_citation(md_text, citation, marker="[^cite1]"):
    if not citation:
        return md_text
    if re.search(r"\[\^[^\]]+\]", md_text):
        # already has authored footnotes; do not inject a duplicate
        return md_text
    lines = md_text.splitlines()
    target = None
    for i, line in enumerate(lines):
        s = line.strip()
        if not s or s.startswith(("#", ">", "<!--", "|", "```")):
            continue
        if len(re.sub(r"[`*_\[\]()]", "", s)) >= 20:
            target = i
            break
    if target is None:
        for i, line in enumerate(lines):
            if line.strip() and not line.startswith("#"):
                target = i
                break
    if target is None:
        return md_text + f"\n\n{marker}\n\n{marker}: {citation}\n"
    lines[target] = lines[target].rstrip() + marker
    return "\n".join(lines) + f"\n\n{marker}: {citation}\n"


def convert(md_path, docx_path, image="", citation="", reference_doc=""):
    md_path = os.path.abspath(md_path)
    docx_path = os.path.abspath(docx_path)
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    pandoc = shutil.which("pandoc")
    if not pandoc:
        return {"success": False, "error": "pandoc not found"}

    # 样式基底：默认用技能内置的中国法学样式集（若存在）
    if not reference_doc:
        default_tpl = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                   "assets", "templates", "中国法学_style.docx")
        reference_doc = default_tpl if os.path.exists(default_tpl) else ""

    text = open(md_path, encoding="utf-8").read()
    text = _inject_citation(text, citation)
    tmp_md = docx_path + ".tmp.md"
    open(tmp_md, "w", encoding="utf-8", newline="\n").write(text)
    cmd = [pandoc, tmp_md, "-o", docx_path, "--from=markdown+footnotes", "--standalone",
           "--resource-path", os.path.dirname(md_path)]
    if reference_doc:
        cmd += ["--reference-doc", reference_doc]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    try:
        os.remove(tmp_md)
    except OSError:
        pass
    if r.returncode != 0:
        return {"success": False, "error": (r.stderr or "")[:500]}
    # 清空模板自带的页眉文字（模板的样例标题不属于本报告）；保留页码页脚
    try:
        from docx import Document
        _doc = Document(docx_path)
        for sec in _doc.sections:
            for p in sec.header.paragraphs:
                for run in p.runs:
                    run.text = ''
            for tbl in sec.header.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        cell.text = ''
        _doc.save(docx_path)
    except Exception:
        pass
    if image and os.path.exists(image):
        try:
            from docx import Document
            from docx.shared import Inches
            doc = Document(docx_path)
            doc.add_page_break()
            doc.add_heading("论证图", level=1)
            doc.add_picture(image, width=Inches(6.5))
            doc.save(docx_path)
        except Exception as e:
            return {"success": True, "docx": docx_path, "warning": f"image insert failed: {e}"}
    return {"success": True, "docx": docx_path, "size_kb": round(os.path.getsize(docx_path) / 1024, 2),
            "footnote_injected": bool(citation), "reference_doc": os.path.basename(reference_doc) if reference_doc else ""}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", default="")
    ap.add_argument("--image", default="")
    ap.add_argument("--citations", default="")
    args = ap.parse_args()
    out = args.out or os.path.splitext(args.md)[0] + ".docx"
    citation = ""
    if args.citations and os.path.exists(args.citations):
        data = json.load(open(args.citations, encoding="utf-8"))
        citation = data.get("citation_law_footnote", "")
    result = convert(args.md, out, args.image, citation)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("success") else 1


if __name__ == "__main__":
    sys.exit(main())
