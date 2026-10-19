#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Deterministic remainder of the anatomy flow.

The agent authors the claim graph (semantic); this script runs everything
deterministic after that:

  PDF -> MD/paragraphs -> metadata -> citations
       -> [graph given] validate + mermaid/render -> report.md -> report.docx
       -> data package

Usage:
  python pipeline.py --pdf paper.pdf --out-dir out/ao --id ao --title "标题" --page 92 --graph out/ao/ao.claim-graph.json
"""
import argparse
import json
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


def run(cmd):
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(f"command failed: {' '.join(cmd)}\n{r.stderr[:1200]}")
    return r.stdout


def _docx_to_artifacts(docx_path, out, doc_id, title):
    """DOCX → <id>.md + <id>.paragraphs.json + <id>.structure.json（与 pdf_to_md 同契约）。

    段落与表格逐块提取；表格行同时进入 md 与 paragraphs（quote 可从表格行锚定）。
    """
    import re as _re
    from docx import Document
    d = Document(docx_path)
    parts, paras, headings = [], [], []
    pid = 0
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        pid += 1
        parts.append(t)
        paras.append({"pid": "P%04d" % pid, "page": 1, "text": t})
        if t[:2] in ("第一", "第二", "第三", "第四", "第五") or _re.match(r"^[（(一二三四五六七八九十]+[）)]", t):
            headings.append({"heading": t[:40], "page": 1, "level": 1})
    for tb in d.tables:
        for row in tb.rows:
            line = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
            if not line.strip(" |"):
                continue
            pid += 1
            parts.append(line)
            paras.append({"pid": "P%04d" % pid, "page": 1, "text": line})
    md_text = "\n\n".join(parts)
    with open(os.path.join(out, doc_id + ".md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(md_text)
    meta = {"id": doc_id, "title": title, "pdf": os.path.abspath(docx_path)}
    with open(os.path.join(out, doc_id + ".paragraphs.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump({**meta, "sections": headings, "paragraphs": paras}, f, ensure_ascii=False, indent=1)
    with open(os.path.join(out, doc_id + ".structure.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump({**meta, "abstract": "", "keywords": "", "sections": headings}, f, ensure_ascii=False, indent=1)
    return f"wrote {doc_id}.md + paragraphs + structure (docx, {len(paras)} 块)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--id", default="")
    ap.add_argument("--title", default="")
    ap.add_argument("--page", default="")
    ap.add_argument("--page-start", default="")
    ap.add_argument("--page-end", default="")
    ap.add_argument("--quote-type", default="paraphrase", choices=["direct", "paraphrase", ""])
    ap.add_argument("--graph", default="")
    ap.add_argument("--reading", action="store_true", help="同时生成阅读/复现提示词任务包（<id>.reading/）")
    ap.add_argument("--skip-docx", action="store_true")
    args = ap.parse_args()

    out = os.path.abspath(args.out_dir)
    os.makedirs(out, exist_ok=True)
    doc_id = args.id or os.path.splitext(os.path.basename(args.pdf))[0]
    py = sys.executable

    # 1. PDF/DOCX -> MD + paragraphs + structure
    if os.path.abspath(args.pdf).lower().endswith(".docx"):
        print(_docx_to_artifacts(args.pdf, out, doc_id, args.title))
    else:
        cmd = [py, os.path.join(HERE, "pdf_to_md.py"), "--pdf", os.path.abspath(args.pdf),
               "--out-dir", out, "--id", doc_id]
        if args.title:
            cmd += ["--title", args.title]
        print(run(cmd).strip())
    md = os.path.join(out, doc_id + ".md")

    # 1.5 阅读/复现提示词任务包（可选）
    if args.reading:
        struct = os.path.join(out, doc_id + ".structure.json")
        rcmd = [py, os.path.join(HERE, "gen_reading_prompts.py"), "--out-dir", out,
                "--id", doc_id, "--title", args.title or doc_id]
        if os.path.exists(struct):
            rcmd += ["--structure", struct]
        print(run(rcmd).strip())

    # 2. metadata
    meta_path = os.path.join(out, doc_id + ".metadata.json")
    meta_cmd = [py, os.path.join(HERE, "legal_metadata.py"), "--md", md,
                "--pdf", os.path.abspath(args.pdf), "--out", meta_path]
    if args.page_start:
        meta_cmd += ["--page-start", args.page_start]
    if args.page_end:
        meta_cmd += ["--page-end", args.page_end]
    print(run(meta_cmd).strip())

    # 3. citations (cite_self + references)
    cit_path = os.path.join(out, doc_id + ".citations.json")
    print(run([py, os.path.join(HERE, "legal_citations.py"), "--metadata", meta_path,
               "--page", args.page or "", "--quote-type", args.quote_type,
               "--out", cit_path]).strip())

    graph = os.path.abspath(args.graph) if args.graph else ""
    report_md = ""
    report_docx = ""
    image = ""
    pkg = ""
    mmd_path = ""
    md_path = ""

    if graph and os.path.exists(graph):
        # 4. validate + render
        v = subprocess.run([py, os.path.join(HERE, "validate_claim_graph.py"), graph],
                           capture_output=True, text=True)
        print(v.stdout.strip())
        if v.returncode != 0:
            print("graph validation failed; abort before report/package")
            return v.returncode
        print(run([py, os.path.join(HERE, "render_claim_graph.py"), "--file", graph,
                   "--out-dir", os.path.join(out, "visuals")]).strip())
        # Mermaid（大图首选；纯 Python 生成，不依赖浏览器）
        mmd_path = os.path.join(out, doc_id + ".claim-graph.mmd")
        md_path = os.path.join(out, doc_id + ".claim-graph.md")
        print(run([py, os.path.join(HERE, "mermaid_claim_graph.py"), "--file", graph,
                   "--out", mmd_path, "--md", md_path]).strip())
        image = os.path.join(out, "visuals", doc_id + ".png")
        if not os.path.exists(image):
            image = ""

        # 5. report.md + report.docx
        report_md = os.path.join(out, doc_id + ".report.md")
        print(run([py, os.path.join(HERE, "build_report.py"), "--graph", graph,
                   "--image", image, "--out", report_md]).strip())
        if not args.skip_docx:
            report_docx = os.path.join(out, doc_id + ".report.docx")
            print(run([py, os.path.join(HERE, "md_to_docx.py"), "--md", report_md,
                       "--out", report_docx]).strip())

        # 6. data package
        print(run([py, os.path.join(HERE, "build_data_package.py"), "--graph", graph,
                   "--out-dir", out]).strip())
        pkg = os.path.join(out, "package")

    manifest = {
        "id": doc_id,
        "pdf": os.path.abspath(args.pdf),
        "article_md": md,
        "paragraphs": os.path.join(out, doc_id + ".paragraphs.json"),
        "structure": os.path.join(out, doc_id + ".structure.json"),
        "metadata": meta_path,
        "citations": cit_path,
        "claim_graph": graph,
        "image": image,
        "mermaid_mmd": mmd_path,
        "mermaid_md": md_path,
        "report_md": report_md,
        "report_docx": report_docx,
        "package": pkg,
    }
    mpath = os.path.join(out, doc_id + ".manifest.json")
    json.dump(manifest, open(mpath, "w", encoding="utf-8", newline="\n"), ensure_ascii=False, indent=2)
    print(json.dumps({"status": "success", "manifest": mpath, "result": manifest}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
